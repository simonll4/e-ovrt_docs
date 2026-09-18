#!/usr/bin/env python3
"""Genera el `.docx` de trabajo de las secciones de costos (§14.4 y §17.2) a partir de los
borradores markdown, usando el informe maestro como plantilla de estilos.

Por qué así: el maestro (`E-OVRT-VDP_v1.1_05062026-sin-indice.docx`) define los estilos
reales del informe —`Ttulo1..3` (heading 1..3), `Normal` con sangría de primera línea,
`Descripcin` (caption) con campo SEQ para "Tabla N", tablas con bordes APA y las tablas
anchas en secciones horizontales—. Construir el documento sobre una copia vaciada del
maestro garantiza que al integrarlo no haya que reformatear nada.

Entrada:  los dos borradores de `informe/entregable/desarrollando/` (sólo el bloque entre
          el primer y el segundo separador `---`; las notas editoriales no viajan).
Salida:   `informe/entregable/desarrollando/E-OVRT-VDP_Secciones_14.4_y_17.2_Costos_Asociados_v0.1.docx`

Convenciones del markdown que entiende:
  `### N. Título`  → heading 2 · `#### N.N.N. Título` → heading 3
  `Tabla N` (línea sola) + línea `*Título*` + tabla `|…|` + `Nota. …` → bloque de tabla
  `- **X.** texto` → ítem con viñeta · `1. **X.** texto` → ítem numerado
  `**negrita**`, `*cursiva*`, `` `[confirmar …]` `` → resaltado amarillo, `` `otro` `` → cursiva

Requiere python-docx (`pip install python-docx`). Uso:
    python3 herramientas/generar_docx_costos.py            # escribe el .docx
    python3 herramientas/generar_docx_costos.py --salida X.docx
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu
except ImportError:  # pragma: no cover
    sys.exit("Falta python-docx: pip install python-docx")

RAIZ = Path(__file__).resolve().parents[1]
DESARROLLANDO = RAIZ / "informe" / "entregable" / "desarrollando"
MAESTRO = DESARROLLANDO / "E-OVRT-VDP_v1.1_05062026-sin-indice.docx"
BORRADORES = [
    DESARROLLANDO / "borrador-14-4-costos-asociados.md",
    DESARROLLANDO / "borrador-17-2-costos-asociados.md",
]
SALIDA = DESARROLLANDO / "E-OVRT-VDP_Secciones_14.4_y_17.2_Costos_Asociados_v0.1.docx"

# Geometría del maestro (Letter, márgenes de 1"): ancho útil en dxa
ANCHO_VERTICAL = 12240 - 2 * 1440    # 9360
ANCHO_HORIZONTAL = 15840 - 2 * 1440  # 12960 (las tablas anchas del maestro miden ~12915)
COLUMNAS_PARA_HORIZONTAL = 5         # con 5+ columnas la tabla va en sección horizontal


# ----------------------------------------------------------------------------- markdown

def bloque_informe(md: str) -> str:
    """Devuelve el texto entre el primer y el segundo `---` (lo que va al informe)."""
    partes = md.split("\n---\n")
    if len(partes) < 3:
        raise SystemExit("El borrador no tiene los dos separadores `---` que delimitan el texto del informe")
    return partes[1].strip("\n")


TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def tramos(texto: str) -> list[tuple[str, dict]]:
    """Parte un texto con marcas inline en tramos (texto, {b,i,hl})."""
    out: list[tuple[str, dict]] = []
    for pieza in TOKEN.split(texto):
        if not pieza:
            continue
        if pieza.startswith("**") and pieza.endswith("**"):
            out.append((pieza[2:-2], {"b": True}))
        elif pieza.startswith("`") and pieza.endswith("`"):
            cuerpo = pieza[1:-1]
            if cuerpo.startswith("["):
                out.append((cuerpo, {"hl": True}))
            else:
                out.append((cuerpo, {"i": True}))
        elif pieza.startswith("*") and pieza.endswith("*") and len(pieza) > 2:
            out.append((pieza[1:-1], {"i": True}))
        else:
            out.append((pieza, {}))
    return out


def parsear(md: str) -> list[dict]:
    """Convierte el bloque markdown en una lista de elementos semánticos."""
    lineas = md.split("\n")
    elems: list[dict] = []
    i = 0
    while i < len(lineas):
        s = lineas[i].rstrip()
        if not s.strip():
            i += 1
            continue
        if s.startswith("#### "):
            elems.append({"tipo": "h3", "texto": s[5:].strip()})
        elif s.startswith("### "):
            elems.append({"tipo": "h2", "texto": s[4:].strip()})
        elif re.fullmatch(r"Tabla \d+", s.strip()):
            numero = s.strip().split()[1]
            i += 1
            while i < len(lineas) and not lineas[i].strip():
                i += 1
            titulo = lineas[i].strip().strip("*")
            i += 1
            while i < len(lineas) and not lineas[i].strip():
                i += 1
            filas: list[list[str]] = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            while i < len(lineas) and not lineas[i].strip():
                i += 1
            nota = None
            if i < len(lineas) and lineas[i].strip().startswith("Nota."):
                nota = lineas[i].strip()[len("Nota."):].strip()
            else:
                i -= 1
            elems.append({"tipo": "tabla", "numero": numero, "titulo": titulo, "filas": filas, "nota": nota})
        elif s.startswith("- "):
            elems.append({"tipo": "vineta", "texto": s[2:].strip()})
        elif re.match(r"^\d+\. ", s):
            n, resto = s.split(". ", 1)
            elems.append({"tipo": "numerado", "n": n, "texto": resto.strip()})
        else:
            elems.append({"tipo": "p", "texto": s.strip()})
        i += 1
    return elems


# --------------------------------------------------------------------------------- docx

def vaciar_cuerpo(doc: Document) -> None:
    """Quita todo el contenido del maestro y deja sólo el sectPr final."""
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    # Sin control de cambios activo: es un documento de trabajo limpio.
    settings = doc.settings.element
    for tag in ("w:trackRevisions",):
        for el in settings.findall(qn(tag)):
            settings.remove(el)
    # Las imágenes del maestro ya no se referencian: soltar sus relaciones para que no
    # viajen en el paquete. Las fuentes embebidas del maestro (`word/fonts/*.odttf`, ~6 MB)
    # sí se conservan a propósito: son las que hacen que el documento se vea igual que el
    # informe, y los demás `.docx` por sección también las llevan.
    for rId, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image") or rel.reltype.endswith("/hyperlink"):
            doc.part.drop_rel(rId)


def agregar_run(p, texto: str, b=False, i=False, hl=False, sz: int | None = None):
    r = p.add_run(texto)
    if b:
        r.bold = True
    if i:
        r.italic = True
    rPr = r._r.get_or_add_rPr()
    if hl:
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rPr.append(h)
    if sz:
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag)
            e.set(qn("w:val"), str(sz))
            rPr.append(e)
    return r


def agregar_tramos(p, texto: str, sz: int | None = None, negrita_global=False):
    for t, f in tramos(texto):
        agregar_run(p, t, b=f.get("b", False) or negrita_global, i=f.get("i", False), hl=f.get("hl", False), sz=sz)


def parrafo_cuerpo(doc, texto: str):
    """Párrafo Normal con sangría de primera línea de 0,5" (como el maestro)."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Emu(720 * 635)  # 720 dxa
    agregar_tramos(p, texto)
    return p


def parrafo_lista(doc, texto: str, marcador: str):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Emu(720 * 635)
    pf.first_line_indent = Emu(-360 * 635)
    agregar_run(p, f"{marcador}\t")
    agregar_tramos(p, texto)
    return p


def titulo(doc, texto: str, nivel: int):
    p = doc.add_paragraph(style=doc.styles[f"Heading {nivel}"])
    p.add_run(texto)
    return p


def caption_tabla(doc, numero: str):
    """'Tabla N' con estilo caption y campo SEQ (número cacheado), como el maestro."""
    p = doc.add_paragraph(style=doc.styles["Caption"])
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    for tag in ("w:i", "w:iCs"):
        e = OxmlElement(tag)
        e.set(qn("w:val"), "0")
        rPr.append(e)
    pPr.append(rPr)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Tabla ")
    r.italic = False

    def fld(tipo):
        run = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tipo)
        run.append(fc)
        return run

    p._p.append(fld("begin"))
    run = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " SEQ Tabla \\* ARABIC "
    run.append(it)
    p._p.append(run)
    p._p.append(fld("separate"))
    rn = p.add_run(numero)
    rn.italic = False
    p._p.append(fld("end"))
    return p


def titulo_tabla(doc, texto: str):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    agregar_run(p, texto, i=True)
    return p


def nota_tabla(doc, texto: str):
    p = doc.add_paragraph()
    agregar_run(p, "Nota", b=True, i=True)
    agregar_run(p, ". ", i=True)
    agregar_tramos(p, texto)
    return p


def _borde(tc_borders, lado, val, sz=None):
    e = OxmlElement(f"w:{lado}")
    e.set(qn("w:val"), val)
    if val != "nil":
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
    tc_borders.append(e)


def anchos_columnas(filas: list[list[str]], total: int) -> list[int]:
    n = len(filas[0])
    largo = [max(len(re.sub(r"[*`]", "", f[c])) if c < len(f) else 0 for f in filas) for c in range(n)]
    peso = [max(10, min(l, 60)) ** 0.75 for l in largo]  # comprime las columnas muy largas
    suma = sum(peso)
    anchos = [int(total * w / suma) for w in peso]
    anchos[-1] += total - sum(anchos)
    return anchos


def tabla(doc, filas: list[list[str]], total: int):
    """Tabla con la apariencia de las del maestro: bordes finos, cabecera con regla gruesa,
    celdas sin bordes laterales, 12 pt, interlineado doble, diseño fijo, centrada."""
    n = len(filas[0])
    t = doc.add_table(rows=len(filas), cols=n)
    tbl = t._tbl
    tblPr = tbl.tblPr
    for el in list(tblPr):
        if el.tag in (qn("w:tblStyle"), qn("w:tblW"), qn("w:tblLook")):
            tblPr.remove(el)
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa"); tblPr.append(tblW)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
    tblInd = OxmlElement("w:tblInd"); tblInd.set(qn("w:w"), "0"); tblInd.set(qn("w:type"), "dxa"); tblPr.append(tblInd)
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _borde(borders, lado, "single", 4)
    tblPr.append(borders)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    look = OxmlElement("w:tblLook"); look.set(qn("w:val"), "0400")
    for k, v in (("firstRow", "0"), ("lastRow", "0"), ("firstColumn", "0"), ("lastColumn", "0"), ("noHBand", "0"), ("noVBand", "1")):
        look.set(qn(f"w:{k}"), v)
    tblPr.append(look)

    anchos = anchos_columnas(filas, total)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(w))

    ultima = len(filas) - 1
    for ri, fila in enumerate(filas):
        tr = t.rows[ri]._tr
        trPr = tr.get_or_add_trPr()
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        for ci in range(n):
            celda = t.cell(ri, ci)
            tcPr = celda._tc.get_or_add_tcPr()
            # python-docx ya crea un tcW por celda: reutilizarlo, nunca duplicarlo
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); tcPr.insert(0, tcW)
            tcW.set(qn("w:w"), str(anchos[ci])); tcW.set(qn("w:type"), "dxa")
            tb = OxmlElement("w:tcBorders")
            if ri == 0:
                _borde(tb, "top", "single", 12); _borde(tb, "left", "nil"); _borde(tb, "right", "nil")
            else:
                _borde(tb, "left", "nil"); _borde(tb, "right", "nil")
                if ri != ultima:
                    _borde(tb, "bottom", "nil")
            tcPr.append(tb)
            mar = OxmlElement("w:tcMar")
            for lado in ("top", "bottom"):
                m = OxmlElement(f"w:{lado}"); m.set(qn("w:w"), "80"); m.set(qn("w:type"), "dxa"); mar.append(m)
            tcPr.append(mar)
            p = celda.paragraphs[0]
            p.paragraph_format.line_spacing = 2.0
            texto = fila[ci] if ci < len(fila) else ""
            agregar_tramos(p, texto, sz=24, negrita_global=(ri == 0))
    return t


def nueva_seccion(doc, horizontal: bool):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Emu(15840 * 635), Emu(12240 * 635)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Emu(12240 * 635), Emu(15840 * 635)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Emu(1440 * 635))
    return sec


def volcar(doc, elems: list[dict]) -> None:
    for e in elems:
        if e["tipo"] == "h2":
            titulo(doc, e["texto"], 2)
        elif e["tipo"] == "h3":
            titulo(doc, e["texto"], 3)
        elif e["tipo"] == "p":
            parrafo_cuerpo(doc, e["texto"])
        elif e["tipo"] == "vineta":
            parrafo_lista(doc, e["texto"], "•")
        elif e["tipo"] == "numerado":
            parrafo_lista(doc, e["texto"], f"{e['n']}.")
        elif e["tipo"] == "tabla":
            horizontal = len(e["filas"][0]) >= COLUMNAS_PARA_HORIZONTAL
            if horizontal:
                nueva_seccion(doc, horizontal=True)
            caption_tabla(doc, e["numero"])
            titulo_tabla(doc, e["titulo"])
            tabla(doc, e["filas"], ANCHO_HORIZONTAL if horizontal else ANCHO_VERTICAL)
            if e["nota"]:
                nota_tabla(doc, e["nota"])
            if horizontal:
                nueva_seccion(doc, horizontal=False)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--salida", type=Path, default=SALIDA)
    ap.add_argument("--maestro", type=Path, default=MAESTRO)
    args = ap.parse_args(argv)

    if not args.maestro.exists():
        sys.exit(f"No encuentro el maestro: {args.maestro}")
    doc = Document(str(args.maestro))
    vaciar_cuerpo(doc)

    for k, ruta in enumerate(BORRADORES):
        elems = parsear(bloque_informe(ruta.read_text(encoding="utf-8")))
        if k > 0:
            doc.add_page_break()
        volcar(doc, elems)

    doc.core_properties.title = "E-OVRT-VDP — Secciones 14.4 y 17.2 (Costos asociados) — documento de trabajo v0.1"
    doc.save(str(args.salida))
    print(f"OK → {args.salida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
